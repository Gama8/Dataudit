"""Exportación de papeles de trabajo a Excel con la trazabilidad que exige la documentación de auditoría (NIA 230)."""
from __future__ import annotations

import io
import re

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = "2B4460"
AGUAMARINA = "49C1C3"
_ENCABEZADO = PatternFill("solid", fgColor=AZUL)
_LINEA = Border(bottom=Side(style="thin", color="BFBFBF"))


def _nombre_hoja(nombre: str, usados: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", nombre).strip()[:31] or "Hoja"
    candidato, i = base, 2
    while candidato.lower() in usados:
        sufijo = f" ({i})"
        candidato, i = base[:31 - len(sufijo)] + sufijo, i + 1
    usados.add(candidato.lower())
    return candidato


def _preparar(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for columna in df.columns:
        serie = df[columna]
        if isinstance(serie.dtype, (pd.CategoricalDtype, pd.PeriodDtype)) or str(serie.dtype) == "string":
            df[columna] = serie.astype(object).where(serie.notna(), None).map(lambda v: v if v is None else str(v))
        elif pd.api.types.is_datetime64_any_dtype(serie) and getattr(serie.dt, "tz", None):
            df[columna] = serie.dt.tz_localize(None)
    df.columns = [str(c) for c in df.columns]
    return df


def papel_de_trabajo(titulo: str, encabezado: dict[str, str], hojas: dict[str, pd.DataFrame],
                     resumen: str = "", conclusion: str = "") -> bytes:
    salida = io.BytesIO()
    usados: set[str] = {"resumen"}
    with pd.ExcelWriter(salida, engine="openpyxl") as escritor:
        libro = escritor.book
        hoja = libro.create_sheet("Resumen")
        hoja["A1"] = "Dataudit · Papel de trabajo"
        hoja["A1"].font = Font(size=10, color=AGUAMARINA, bold=True)
        hoja["A2"] = titulo
        hoja["A2"].font = Font(size=16, color=AZUL, bold=True)
        fila = 4
        for clave, valor in encabezado.items():
            hoja.cell(fila, 1, clave).font = Font(bold=True, color=AZUL)
            celda = hoja.cell(fila, 2, valor)
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            hoja.cell(fila, 1).border = hoja.cell(fila, 2).border = _LINEA
            fila += 1
        for etiqueta, texto in (("Resultado", resumen), ("Conclusión del auditor", conclusion or "")):
            fila += 1
            hoja.cell(fila, 1, etiqueta).font = Font(bold=True, color=AZUL)
            celda = hoja.cell(fila, 2, texto)
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            hoja.row_dimensions[fila].height = max(30, 15 * (len(texto) // 90 + 1))
        fila += 2
        for rol in ("Preparado por", "Revisado por"):
            hoja.cell(fila, 1, rol).font = Font(bold=True, color=AZUL)
            hoja.cell(fila, 2, "Nombre: ______________________   Firma: ______________   Fecha: __________")
            fila += 1
        hoja.column_dimensions["A"].width = 30
        hoja.column_dimensions["B"].width = 100

        for nombre, df in hojas.items():
            nombre_hoja = _nombre_hoja(nombre, usados)
            datos = _preparar(df)
            if datos.empty:
                datos = pd.DataFrame({"Resultado": ["Sin registros"]})
            datos.to_excel(escritor, sheet_name=nombre_hoja, index=False)
            ws = escritor.sheets[nombre_hoja]
            for celda in ws[1]:
                celda.fill = _ENCABEZADO
                celda.font = Font(bold=True, color="FFFFFF")
                celda.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, columna in enumerate(datos.columns, start=1):
                muestra = datos[columna].head(200).map(lambda v: len(str(v)))
                ancho = min(60, max(len(str(columna)), int(muestra.max()) if len(muestra) else 0) + 2)
                ws.column_dimensions[get_column_letter(i)].width = max(10, ancho)
                if pd.api.types.is_datetime64_any_dtype(datos[columna]):
                    for celda in ws[get_column_letter(i)][1:]:
                        celda.number_format = "yyyy-mm-dd"
                elif pd.api.types.is_float_dtype(datos[columna]):
                    for celda in ws[get_column_letter(i)][1:]:
                        celda.number_format = "#,##0.00"
        if "Sheet" in libro.sheetnames and len(libro.sheetnames) > 1:
            del libro["Sheet"]
        libro.move_sheet("Resumen", offset=-libro.index(libro["Resumen"]))
        libro.active = 0
    return salida.getvalue()
